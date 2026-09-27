"""Выгрузка колоды в HTML (`mimeo/export/html.py`).

Колоды для проверки собирает `python-pptx` — независимый писатель: так тест
ловит и то, что наш собственный сборщик никогда не пишет. Проверяется то, что
видит зритель: весь текст на месте и остаётся текстом, картинка встроена,
таблица — таблица, у диаграммы есть подписи, положение фигуры — в долях
холста, повтор даёт тот же байт.
"""

from __future__ import annotations

import html
import re
import struct
import zipfile
import zlib

import pytest

pytest.importorskip("pptx")

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from mimeo.export import render_html, to_html

TITLE = "Проверка выгрузки"
BULLETS = ["Первый пункт списка", "Второй пункт", "Вложенный пункт"]
CELLS = [["Показатель", "2025", "2026"], ["Выручка", "12", "17"], ["Клиенты", "340", "410"]]
CELL_TEXTS = [text for row in CELLS for text in row]
SHAPE_TEXTS = ["Повёрнутый текст", "В группе"]
CATEGORIES = ["Январь", "Февраль", "Март"]
VALUES = (11.0, 23.0, 31.0)
PIE = ("Север", "Юг", "Запад")
PIE_VALUES = (45.0, 35.0, 20.0)


def _png(width: int = 4, height: int = 3) -> bytes:
    """Маленький PNG без сторонних библиотек."""
    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))
    raw = b"".join(b"\x00" + bytes((200, 40, 40) * width) for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _deck(path, png_path) -> dict:
    """Колода на пять слайдов; возвращает то, что тесту надо знать о ней."""
    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    info: dict = {"W": W, "H": H}

    # 1. Заголовок и список с уровнями.
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = TITLE
    body = s.placeholders[1].text_frame
    body.text = BULLETS[0]
    for i, text in enumerate(BULLETS[1:], 1):
        p = body.add_paragraph()
        p.text = text
        p.level = 1 if i == 2 else 0

    # 2. Картинка, повёрнутая фигура с текстом, группа, свободная форма.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(png_path), Emu(W // 10), Emu(H // 10), Emu(W // 5), Emu(H // 5))
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(W // 4), Emu(H // 2), Emu(W // 5), Emu(H // 6))
    box.rotation = 30
    box.text_frame.text = "Повёрнутый текст"
    group = s.shapes.add_group_shape()
    group.shapes.add_shape(MSO_SHAPE.OVAL, Emu(W // 2), Emu(H // 10), Emu(W // 10), Emu(H // 10))
    inner = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(W * 7 // 10), Emu(H // 10),
                                   Emu(W // 10), Emu(H // 10))
    inner.text_frame.text = "В группе"
    ff = s.shapes.build_freeform(Emu(W * 6 // 10), Emu(H * 6 // 10), scale=1.0)
    ff.add_line_segments([(W * 8 // 10, H * 6 // 10), (W * 7 // 10, H * 8 // 10)], close=True)
    ff.convert_to_shape()

    # 3. Таблица.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    frame = s.shapes.add_table(3, 3, Emu(W // 10), Emu(H // 10), Emu(W * 8 // 10), Emu(H // 3))
    for r, row in enumerate(CELLS):
        for c, text in enumerate(row):
            frame.table.cell(r, c).text = text

    # 4. Две диаграммы с подписями значений.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = CATEGORIES
    data.add_series("Продажи", VALUES)
    col = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(0), Emu(0), Emu(W // 2), Emu(H // 2),
                             data).chart
    col.plots[0].has_data_labels = True
    col.plots[0].data_labels.show_value = True
    pie_data = CategoryChartData()
    pie_data.categories = PIE
    pie_data.add_series("Доли", PIE_VALUES)
    pie = s.shapes.add_chart(XL_CHART_TYPE.PIE, Emu(W // 2), Emu(0), Emu(W // 2), Emu(H // 2),
                             pie_data).chart
    pie.has_legend = True
    pie.plots[0].has_data_labels = True
    pie.plots[0].data_labels.show_value = True

    # 5. Скрытый слайд: в выгрузку не попадает.
    hidden = prs.slides.add_slide(prs.slide_layouts[6])
    hidden.shapes.add_textbox(Emu(0), Emu(0), Emu(W // 2), Emu(H // 4)).text_frame.text = "Скрыто"
    hidden._element.set("show", "0")

    prs.save(str(path))
    return info


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    folder = tmp_path_factory.mktemp("export")
    png = folder / "pic.png"
    png.write_bytes(_png())
    deck = folder / "deck.pptx"
    info = _deck(deck, png)
    text, report = to_html(str(deck))
    return deck, info, text, report


def _visible(text: str) -> str:
    """Текст страницы без разметки — то, что увидит и найдёт поиском зритель."""
    return html.unescape(re.sub(r"<[^>]+>", "", text))


def test_every_run_is_text(built):
    _, _, text, report = built
    seen = _visible(text)
    for piece in [TITLE, *BULLETS, *SHAPE_TEXTS, *CELL_TEXTS]:
        assert piece in seen, piece
    assert "Скрыто" not in seen
    # Счётчик считает знаки текста слайдов, выведенные текстом (без скрытого слайда).
    expected = sum(len(x) for x in [TITLE, *BULLETS, *SHAPE_TEXTS, *CELL_TEXTS])
    assert report.text_chars == expected
    assert report.slides == 4
    assert any(u.startswith("hidden slide ×1") for u in report.unsupported)


def test_document_is_self_contained(built):
    _, _, text, _ = built
    assert text.startswith("<!doctype html>")
    assert '<html lang="ru">' in text
    assert "data:image/png;base64," in text
    # Ничего не грузится снаружи: ни стилей, ни скриптов, ни шрифтов, ни картинок.
    assert "<link" not in text
    assert not re.search(r"""(src|href)\s*=\s*["']?https?:""", text)
    assert "url(http" not in text
    assert text.count('<section class="slide"') == 4


def test_picture_is_embedded_once_and_counted(built):
    _, _, text, report = built
    assert report.images == 1
    assert text.count("data:image/png;base64,") == 1


def test_position_is_percent_of_canvas(built):
    """Фигура с левым краем на четверти ширины стоит на `left:25%` — и так при
    любой ширине окна."""
    _, _, text, _ = built
    assert re.search(r'class="s" style="left:25%;top:50%;width:20%;height:16\.66\d*%;'
                     r'transform:rotate\(30deg\)"', text)


def test_group_child_keeps_its_place(built):
    """Группа python-pptx пишет систему координат детей равной своей: ребёнок
    на 70% ширины и 10% высоты там и остаётся."""
    _, _, text, _ = built
    assert re.search(r'class="s" style="left:70%;top:10%;width:10%;height:10%"', text)


def test_custom_geometry_is_svg_path(built):
    """Свободная форма (`a:custGeom`) — путь SVG в пунктах рамки фигуры:
    треугольник 144×108 пт из точек колоды."""
    _, _, text, _ = built
    slide2 = text.split('id="slide-2"')[1].split('id="slide-3"')[0]
    assert '<path d="M0 0L144 0L72 108Z"' in slide2
    # Заливка по ссылке стиля фигуры на тему — градиент python-pptx по умолчанию.
    assert "<linearGradient" in slide2


def test_table_is_a_table(built):
    _, _, text, _ = built
    slide3 = text.split('id="slide-3"')[1].split('id="slide-4"')[0]
    assert "<table" in slide3
    cells = re.findall(r"<td[^>]*>(.*?)</td>", slide3)
    assert [_visible(c) for c in cells] == CELL_TEXTS
    # Стиль по умолчанию (Medium Style 2): шапка жирная.
    first_row = slide3.split("</tr>")[0]
    assert "font-weight:700" in first_row


@pytest.mark.parametrize("name, drawn", [("Table_0", False), ("Medium Style 2 - Accent 1", True)])
def test_only_builtin_table_styles_are_drawn(built, tmp_path, name, drawn):
    """PowerPoint рисует только свои встроенные стили таблиц; чужое описание
    в `tableStyles.xml` (Google Slides пишет «Table_0») он не рисует вовсе —
    рендеры шаблонов VK Education и VK WorkSpace 26 сентября."""
    deck, _, _, _ = built
    sid = "{11111111-2222-3333-4444-555555555555}"
    style = (f'<a:tblStyleLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
             f'def="{sid}"><a:tblStyle styleId="{sid}" styleName="{name}"><a:wholeTbl><a:tcStyle>'
             f'<a:fill><a:solidFill><a:srgbClr val="FF00AA"/></a:solidFill></a:fill></a:tcStyle>'
             f'</a:wholeTbl></a:tblStyle></a:tblStyleLst>').encode()
    patched = tmp_path / "styles.pptx"
    with zipfile.ZipFile(deck) as src, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "ppt/slides/slide3.xml":
                data = re.sub(rb"<a:tableStyleId>[^<]*</a:tableStyleId>",
                              b"<a:tableStyleId>" + sid.encode() + b"</a:tableStyleId>", data)
            elif item.filename == "ppt/tableStyles.xml":
                data = style
            dst.writestr(item, data)
    text, report = to_html(str(patched))
    assert ("background:#FF00AA" in text) is drawn
    assert not any(u.startswith("table style") for u in report.unsupported)


def test_charts_are_svg_with_labels(built):
    _, _, text, report = built
    slide4 = text.split('id="slide-4"')[1]
    charts = re.findall(r'<svg class="ch".*?</svg>', slide4)
    assert len(charts) == 2
    column, pie = (_visible(c) for c in charts)
    for name in CATEGORIES:
        assert name in column
    for value in ("11", "23", "31"):
        assert value in column
    for name in PIE:
        assert name in pie
    for value in ("45", "35", "20"):
        assert value in pie
    assert not any(u.startswith("chart") for u in report.unsupported)


def test_output_is_deterministic(built, tmp_path):
    deck, _, text, _ = built
    again, _ = to_html(str(deck))
    assert again == text
    out = tmp_path / "deck.html"
    report = render_html(str(deck), str(out))
    raw = out.read_bytes()
    assert raw == text.encode("utf-8")
    assert b"\r\n" not in raw
    assert report.output == str(out)


def test_emf_picture_is_skipped_and_counted(built, tmp_path):
    """EMF браузер не покажет: картинку пропускаем и называем в отчёте, а
    выгрузка не падает."""
    deck, _, _, _ = built
    emf = struct.pack("<II", 1, 88) + b"\x00" * 32 + b" EMF" + b"\x00" * 44
    patched = tmp_path / "emf.pptx"
    with zipfile.ZipFile(deck) as src, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("ppt/media/"):
                data = emf
            dst.writestr(item, data)
    text, report = to_html(str(patched))
    assert "EMF picture ×1" in report.unsupported
    assert report.images == 0
    assert "data:image" not in text
    assert TITLE in _visible(text)


_DGM_FRAME = (
    '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="90" name="Diagram"/><p:cNvGraphicFramePr/>'
    '<p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="914400" y="914400"/>'
    '<a:ext cx="3657600" cy="1828800"/></p:xfrm><a:graphic><a:graphicData '
    'uri="http://schemas.openxmlformats.org/drawingml/2006/diagram"><dgm:relIds '
    'xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" r:dm="rId90" '
    'r:lo="rId92" r:qs="rId93" r:cs="rId94"/></a:graphicData></a:graphic></p:graphicFrame>')
_DGM_DATA = (
    '<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:dsp="http://schemas.microsoft.com/office/drawing/2008/diagram"><dgm:ptLst/>'
    '<dgm:extLst><a:ext uri="http://schemas.microsoft.com/office/drawing/2008/diagram">'
    '<dsp:dataModelExt relId="rId91" minVer="http://schemas.openxmlformats.org/drawingml/2006/diagram"/>'
    '</a:ext></dgm:extLst></dgm:dataModel>')
_DGM_DRAWING = (
    '<dsp:drawing xmlns:dsp="http://schemas.microsoft.com/office/drawing/2008/diagram" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dsp:spTree>'
    '<dsp:nvGrpSpPr><dsp:cNvPr id="0" name=""/><dsp:cNvGrpSpPr/></dsp:nvGrpSpPr><dsp:grpSpPr/>'
    '<dsp:sp modelId="{1}"><dsp:nvSpPr><dsp:cNvPr id="0" name=""/><dsp:cNvSpPr/></dsp:nvSpPr>'
    '<dsp:spPr><a:xfrm><a:off x="914400" y="0"/><a:ext cx="1828800" cy="914400"/></a:xfrm>'
    '<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="2E75B6"/>'
    '</a:solidFill></dsp:spPr><dsp:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/><a:p>'
    '<a:pPr algn="ctr"/><a:r><a:rPr lang="ru-RU" sz="1400"/><a:t>Шаг SmartArt</a:t></a:r></a:p>'
    '</dsp:txBody><dsp:txXfrm><a:off x="960120" y="45720"/><a:ext cx="1737360" cy="822960"/>'
    '</dsp:txXfrm></dsp:sp></dsp:spTree></dsp:drawing>')


def test_smartart_is_drawn_from_saved_layout(tmp_path):
    """SmartArt рисуется по готовому рисунку PowerPoint (`ppt/diagrams/drawing`):
    фигура стоит в рамке кадра, текст — текстом."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    plain = tmp_path / "plain.pptx"
    prs.save(str(plain))
    rels = ('<Relationship Id="rId90" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/diagramData" Target="../diagrams/data1.xml"/><Relationship Id="rId91" '
            'Type="http://schemas.microsoft.com/office/2007/relationships/diagramDrawing" '
            'Target="../diagrams/drawing1.xml"/>')
    deck = tmp_path / "smartart.pptx"
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                data = data.replace(b"</p:spTree>", _DGM_FRAME.encode("utf-8") + b"</p:spTree>")
            elif item.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace(b"</Relationships>", rels.encode("utf-8") + b"</Relationships>")
            dst.writestr(item, data)
        dst.writestr("ppt/diagrams/data1.xml", _DGM_DATA.encode("utf-8"))
        dst.writestr("ppt/diagrams/drawing1.xml", _DGM_DRAWING.encode("utf-8"))
    text, report = to_html(str(deck))
    assert "Шаг SmartArt" in _visible(text)
    # Кадр в (914400, 914400) плюс фигура в (914400, 0) внутри него: 20% ширины слайда.
    assert re.search(r'class="s" style="left:20%;top:13\.33\d*%;width:20%;height:13\.33\d*%"', text)
    assert "background:#2E75B6" in text
    assert not any("SmartArt" in u for u in report.unsupported)


def test_percent_strings_do_not_break_export(built, tmp_path):
    """Часть генераторов пишет проценты как "50%" (так в шаблоне VK WorkSpace).
    Разбор не падает, а заливка получается."""
    deck, _, _, _ = built
    background = (
        b"<p:bg><p:bgPr><a:gradFill><a:gsLst>"
        b"<a:gs pos=\"0%\"><a:srgbClr val=\"FF0000\"><a:alpha val=\"50%\"/></a:srgbClr></a:gs>"
        b"<a:gs pos=\"100%\"><a:srgbClr val=\"0000FF\"/></a:gs></a:gsLst>"
        b"<a:path path=\"circle\"><a:fillToRect l=\"50%\" t=\"50%\" r=\"50%\" b=\"50%\"/>"
        b"</a:path></a:gradFill><a:effectLst/></p:bgPr></p:bg>")
    patched = tmp_path / "pct.pptx"
    with zipfile.ZipFile(deck) as src, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "ppt/slides/slide2.xml":
                data = re.sub(rb"(<p:cSld[^>]*>)", lambda m: m.group(1) + background, data, count=1)
            dst.writestr(item, data)
    text, report = to_html(str(patched))
    slide2 = text.split('id="slide-2"')[1].split(">")[0]
    assert "radial-gradient(circle farthest-corner at 50% 50%,rgba(255,0,0,0.5) 0%" in slide2
    assert not any(u.startswith("render error") for u in report.unsupported)
