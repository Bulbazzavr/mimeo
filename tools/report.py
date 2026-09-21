"""Порождение таблиц замеров для WORKLOG.

Таблицы в `WORKLOG/*.md` не пишутся руками — они выводятся прогоном по корпусу
образцов. Этот скрипт их и печатает.

    python tools/report.py analyze     разбор шаблонов
    python tools/report.py plan        раскладка контента по шаблонам
    python tools/report.py compose     сборка готовых файлов и три проверки
    python tools/report.py volume      объём колоды: без цели и с целью (`--slides`)
    python tools/report.py llm         цена обращения к модели: сколько вызовов на колоду
                                       и какого размера каждый (`--content путь`)
    python tools/report.py slots       сужённые и пустые слоты (`Z-48`, `Z-49`); PowerPoint не нужен
    python tools/report.py verify      переполнения, заслонения и время настоящим PowerPoint
                                       (`--variants 3` на tz/templates — девять сдаточных колод)
    python tools/report.py prose       что даёт вход: форма, ёмкость слотов, колода
                                       (`--content путь`, по умолчанию прозаический пример;
                                       `--raw` — без сегментации, как было до `Z-25`)

Печатает в стандартный вывод намеренно: файлы в `WORKLOG/` содержат ещё и
разбор наблюдений, написанный руками, и перезаписывать их целиком нельзя.
Новые числа надо посмотреть и осознанно вклеить, а не затереть анализ.

Образцы в `samples/` не коммитятся (они чужие). Скачать заново:

    python tools/fetch_samples.py samples
"""

from __future__ import annotations

import argparse
import glob
import os
import sys
import time
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from mimeo.analyze import analyze_template  # noqa: E402
from mimeo.compose import build, inspect  # noqa: E402
from mimeo.oxml.units import emu_to_inch as inch  # noqa: E402
from mimeo.plan import load_content, load_markdown, plan_deck, rank  # noqa: E402
from mimeo.cli import parse_slides  # noqa: E402
from mimeo.plan.prompt import build_request  # noqa: E402

CONTENT = "examples/content-demo.md"


#: Каталог с шаблонами. Меняется флагом `--templates`: те же таблицы нужны и по
#: нашему корпусу-прокси (`samples/`), и по выданным материалам (`tz/templates/`),
#: и смешивать их в одной таблице нельзя — у них разный статус (`tz/README.md`).
TEMPLATES = "samples"


def samples() -> list[str]:
    found = sorted(glob.glob(os.path.join(TEMPLATES, "*.pptx")))
    if not found:
        sys.exit(
            f"В {TEMPLATES}/ нет ни одного .pptx. Для samples/ скачать: "
            "python tools/fetch_samples.py samples"
        )
    return found


def report_analyze() -> None:
    print("| Файл | Холст, дюймы | Сл. | Мак. | Цветов | Ядро | Типо-ролей | Форм | "
          "Фигур сод./поз. | Паттернов | Слотов | Сек. |")
    print("|---|---|---|---|---|---|---|---|---|---|---|---|")
    rows = []
    for f in samples():
        started = time.perf_counter()
        a = analyze_template(f)
        elapsed = time.perf_counter() - started
        d, p, e = a.design_system, a.patterns, a.design_system.evidence
        rows.append((os.path.basename(f), d, p))
        print(
            f"| `{os.path.basename(f)}` | {inch(d.slide.cx_emu):.1f}×{inch(d.slide.cy_emu):.1f} "
            f"| {d.source.slides} | {d.source.layouts} | {len(d.observed_palette)} "
            f"| {len(d.core_palette)} | {len(d.type_scale)} | {len(d.shapes)} "
            f"| {e.shapes_used}/{e.shapes_positioned} | {len(p.patterns)} "
            f"| {sum(len(x.slots) for x in p.patterns)} | {elapsed:.2f} |"
        )

    print("\n### Поля, выведенные по текстовым блокам\n")
    print("| Файл | Л | В | П | Н | Колонок |")
    print("|---|---|---|---|---|---|")
    for name, d, _ in rows:
        g = d.grid
        print(f"| `{name}` | {inch(g.margin_left_emu):.2f} | {inch(g.margin_top_emu):.2f} "
              f"| {inch(g.margin_right_emu):.2f} | {inch(g.margin_bottom_emu):.2f} "
              f"| {g.columns} |")

    print("\n### Размеры холста\n")
    print("| EMU | Дюймы | Файлов |")
    print("|---|---|---|")
    sizes = Counter((d.slide.cx_emu, d.slide.cy_emu) for _, d, _ in rows)
    for (cx, cy), n in sorted(sizes.items(), key=lambda kv: -kv[1]):
        print(f"| {cx} × {cy} | {inch(cx):.2f} × {inch(cy):.2f} | {n} |")

    print("\n### Кластеризация\n")
    print("| Файл | Слайдов | Паттернов | Слитых групп | Размеры групп |")
    print("|---|---|---|---|---|")
    for name, d, p in rows:
        merged = [len(x.members) for x in p.patterns if len(x.members) > 1]
        print(f"| `{name}` | {d.source.slides} | {len(p.patterns)} | {len(merged)} "
              f"| {merged or '—'} |")


def report_plan() -> None:
    doc = load_content(CONTENT)
    print("| Шаблон | Паттернов | Слайдов | Не размещено | Сверх ёмкости | Предупр. "
          "| Виды раскладок |")
    print("|---|---|---|---|---|---|---|---|")
    clean = 0
    total = 0
    for f in samples():
        total += 1
        a = analyze_template(f)
        p = plan_deck(doc, a.patterns, a.design_system.source.sha256)
        kinds = {x.id: x.kind for x in a.patterns.patterns}
        sequence = " → ".join(kinds.get(s.pattern_id, "?") for s in p.slides)
        over = sum(1 for s in p.slides for fill in s.fills if fill.over_capacity)
        clean += not p.unplaced
        print(f"| `{os.path.basename(f)}` | {len(a.patterns.patterns)} | {len(p.slides)} "
              f"| {len(p.unplaced)} | {over} | {len(p.warnings)} | {sequence} |")
    print(f"\n**Итог: {clean} шаблонов из {total} разложены без потерь.**")


def report_compose() -> None:
    doc = load_content(CONTENT)
    os.makedirs("out/decks", exist_ok=True)
    print("| Шаблон | Слайдов | Подставлено слотов | Унаследовано фигур "
          "| Структурных проблем | Круговая | python-pptx | КБ | Сек |")
    print("|---|---|---|---|---|---|---|---|---|---|")
    bad = 0
    total = 0
    for f in samples():
        total += 1
        name = os.path.basename(f)
        target = f"out/decks/{name}"
        started = time.perf_counter()
        a = analyze_template(f)
        plan = plan_deck(doc, a.patterns, a.design_system.source.sha256)
        report = build(f, plan, a.patterns, target)
        problems = len(inspect(target))

        try:
            circular = f"{analyze_template(target).design_system.source.slides} сл."
        except Exception as exc:  # noqa: BLE001 — это и есть проверка
            circular = f"ПАДЕНИЕ {type(exc).__name__}"
        try:
            from pptx import Presentation

            third = f"{len(Presentation(target).slides)} сл."
        except ImportError:
            third = "нет python-pptx"
        except Exception as exc:  # noqa: BLE001
            third = f"ПАДЕНИЕ {type(exc).__name__}"

        if problems or "ПАДЕНИЕ" in circular or "ПАДЕНИЕ" in third:
            bad += 1
        print(f"| `{name}` | {report.slides} | {report.substituted} "
              f"| {report.inherited_shapes} | {problems} | {circular} | {third} "
              f"| {os.path.getsize(target) // 1024} | {time.perf_counter() - started:.2f} |")
    print(f"\n**Итог: {total - bad} шаблонов из {total} собраны и прошли все три проверки.**")


def _empty_required(plan, patterns) -> int:
    """Обязательные слоты, которым не досталось содержимого."""
    by_id = {p.id: p for p in patterns.patterns}
    empty = 0
    for slide in plan.slides:
        pattern = by_id.get(slide.pattern_id)
        if pattern is None:
            continue
        filled = {f.slot_id for f in slide.fills}
        empty += sum(1 for s in pattern.slots if s.required and s.id not in filled)
    return empty


def report_slots(content: str) -> None:
    """Что стало со слотами: сужённые, пустые, сколько знаков разместилось.

    PowerPoint не нужен — всё считается по плану. Это дешёвая половина приёмки
    `Z-48` и `Z-49`: сужение ёмкости заслонённых слотов (`Slot.occluded`) и его
    цена — обязательные слоты, которым не досталось содержимого.

    **Считает по одному каталогу за раз**, как и остальные стадии: смешивать
    наш корпус-прокси с выданными материалами нельзя, у них разный статус
    (`tz/README.md`). Числа по всем четырнадцати шаблонам — сумма двух прогонов.

    Числа за 21 сентября на `examples/content-mimeo.md`:

    | Прогон | Слотов с ёмкостью | Сужено | Пустых обязательных | Знаков |
    |---|---:|---:|---:|---:|
    | `samples` (11) | 658 | 17 | 191 | 25 458 |
    | `--templates tz/templates` (3) | 999 | 32 | 32 | 7 512 |
    | **сумма, 14 шаблонов** | **1657** | **49** | **223** | **32 970** |

    До `Z-48` пустых обязательных было **217**; рост — цена того, что
    заслонённый слот перестал получать запас на переполнение (`Z-49`).
    """
    doc = load_content(content)
    print(f"Вход: `{content}`\n")
    print("| Шаблон | Слотов с ёмкостью | Сужено | Пустых обязательных | Из них | Знаков |")
    print("|---|---:|---:|---:|---:|---:|")
    cap = narrowed = empty = required = chars = 0
    for f in samples():
        a = analyze_template(f)
        plan = plan_deck(doc, a.patterns, a.design_system.source.sha256)
        c = sum(1 for p in a.patterns.patterns for s in p.slots if s.capacity is not None)
        # Сужённые считаются **среди слотов с ёмкостью**, а не среди всех:
        # у слота под картинку или диаграмму ёмкости нет, сужать нечего, и
        # смешивать эти два счёта — та самая ошибка «два измерителя на разные
        # вопросы» (`CLAUDE.md`). Без фильтра получается 146 вместо 49.
        n = sum(1 for p in a.patterns.patterns for s in p.slots
                if s.capacity is not None and s.occluded)
        e = _empty_required(plan, a.patterns)
        r = sum(1 for p in a.patterns.patterns for s in p.slots if s.required)
        ch = sum(len(fi.text or "") + sum(len(i) for i in (fi.items or ()))
                 for sl in plan.slides for fi in sl.fills)
        cap += c; narrowed += n; empty += e; required += r; chars += ch
        print(f"| `{os.path.basename(f)}` | {c} | {n} | {e} | {r} | {ch} |")
    print(f"| **всего** | **{cap}** | **{narrowed}** | **{empty}** | **{required}** "
          f"| **{chars}** |")
    print("\n«Сужено» — у скольких слотов видимая полоса у́же бокса (`Z-48`). "
          "«Пустых обязательных» — цена этого сужения, задача `Z-49`.")


def report_verify(content: str, slides: str | None, variants: int) -> None:
    """Переполнения, заслонения и время — настоящим PowerPoint.

    **Поднимает приложение на рабочем столе** (`ADR-0013`). Это та самая
    таблица, на которой стоят числа `Z-38`, `Z-47` и `Z-48`, и до 21 сентября
    её приходилось собирать разовым скриптом — то есть числа в документах
    нечем было перепроверить.

    `--variants 3` на `--templates tz/templates` даёт девять сдаточных колод.
    """
    from mimeo.verify import verify_deck                        # noqa: PLC0415

    doc = load_content(content, target=parse_slides(slides))
    os.makedirs("out/verify", exist_ok=True)
    print(f"Вход: `{content}`, вариантов {variants}"
          + (f", объём {slides}" if slides else "") + "\n")
    print("| Шаблон | Вариант | Слайдов | Переполнений до | после | Заслонений "
          "| Пустых | Чем встала | Вся команда, с | Петля, с |")
    print("|---|---:|---:|---:|---:|---:|---:|---|---:|---:|")
    before = after = occl = 0
    for f in samples():
        a = analyze_template(f)
        plans = _variant_plans(doc, a, variants)
        for n, plan in enumerate(plans, 1):
            started = time.perf_counter()
            target = os.path.join("out/verify",
                                  f"{os.path.splitext(os.path.basename(f))[0]}-{n}.pptx")
            built = build(f, plan, a.patterns, target)
            outcome = verify_deck(f, plan, a.patterns, target, built, rounds=3)
            rep = outcome.report
            before += rep.before or 0
            after += rep.after or 0
            occl += len(rep.occluded)
            print(f"| `{os.path.basename(f)}` | {n} | {len(outcome.plan.slides)} "
                  f"| {rep.before} | {rep.after} | {len(rep.occluded)} "
                  f"| {_empty_required(outcome.plan, a.patterns)} | {rep.stopped} "
                  f"| {time.perf_counter() - started:.1f} | {rep.seconds:.1f} |")
    print(f"\n**Итого: переполнений {before} → {after}, заслонений {occl}.** "
          "Заслонения ремонтом не берутся и в «до/после» не входят (`Z-47`).")


def _variant_plans(doc, analysis, variants: int) -> list:
    """Один план или тройка вариантов — тем же путём, каким их строит `build`.

    Важно брать именно `generate` + `select`, а не свою выборку: иначе таблица
    померила бы не то, что уходит в сдачу (`ADR-0020`).
    """
    sha = analysis.design_system.source.sha256
    if variants <= 1:
        return [plan_deck(doc, analysis.patterns, sha)]
    from mimeo.plan.variants import generate, load_policies, select   # noqa: PLC0415

    policies, min_distance, _ = load_policies()
    chosen, _ = select(generate(doc, analysis.patterns, sha, policies, None),
                       variants, min_distance)
    return [v.plan for v in chosen]


def _quartiles(values: list[int]) -> tuple[int, int, int]:
    """Медиана и квартили. Пустой список — не ноль, а отсутствие ответа."""
    if not values:
        raise ValueError("нечего мерить")
    v = sorted(values)
    return v[len(v) // 4], v[len(v) // 2], v[3 * len(v) // 4]


def report_llm(content: str) -> None:
    """Цена обращения к модели: сколько вызовов на колоду и какого размера.

    Обращения ещё нет — считается то, что УШЛО БЫ, если звать модель на каждую
    секцию. Нужно, чтобы решать про кэш и про потолок в пять минут из ТЗ, а не
    прикидывать. Числа последнего прогона — `WORKLOG/2026-09-17-llm-client.md`.
    """
    doc = load_content(content)
    print("| Шаблон | Паттернов | Секций | Макс. запрос, зн. | Сумма, зн. |")
    print("|---|---|---|---|---|")
    peaks: list[int] = []
    counts: list[int] = []
    for f in samples():
        a = analyze_template(f)
        by_id = {x.id: x for x in a.patterns.patterns}
        sizes = []
        for section in doc.sections:
            matches = rank(section, a.patterns.patterns)
            if not matches:
                continue
            r = build_request(section, matches, by_id, a.design_system)
            sizes.append(len(r.system) + len(r.user))
        peaks.append(max(sizes) if sizes else 0)
        counts.append(len(sizes))
        print(f"| `{os.path.basename(f)}` | {len(a.patterns.patterns)} | {len(sizes)} "
              f"| {max(sizes) if sizes else 0} | {sum(sizes)} |")
    if peaks:
        print()
        print(f"**Вызовов на колоду {min(counts)}–{max(counts)}, "
              f"самый большой запрос {min(peaks)}–{max(peaks)} знаков.** "
              f"Контент: `{content}`.")


def report_prose(content: str, raw: bool = False) -> None:
    """Замеры по `Z-25`: что вход даёт на корпусе и во что упирается.

    Отдельно от `plan`, потому что меряет не шаблон, а **вход**: одна и та же
    таблица гоняется по прозе и по размеченному контенту, и сравниваются они
    между собой.

    `raw=True` обходит сегментацию и показывает, что было до `Z-25`. Нужен, чтобы
    таблица «до» в `WORKLOG/2026-09-14-prose-input.md` осталась воспроизводимой:
    иначе документ обещал бы числа, которых уже ничем не получить.
    """
    doc = load_markdown(content) if raw else load_content(content)
    mode = "без сегментации, как до `Z-25`" if raw else "с сегментацией"
    print(f"### Форма входа: `{os.path.basename(content)}` — {mode}\n")
    print(f"Разделов {len(doc.sections)}, блоков {doc.block_count}, "
          f"заголовок колоды {'есть' if doc.title else 'нет'}.\n")
    print("| Раздел | Заголовок | Блок | Знаков | Единиц |")
    print("|---|---|---|---|---|")
    for s in doc.sections:
        for b in s.blocks:
            head = (s.heading or "—")[:36]
            print(f"| {s.id} | {head} | {b.kind} | {b.length} | {b.units} |")

    print("\n### Ёмкость слотов корпуса по ролям\n")
    print("| Роль | Слотов | max_chars: кв1 | медиана | кв3 |")
    print("|---|---|---|---|---|")
    by_role: dict[str, list[int]] = {}
    for f in samples():
        for pattern in analyze_template(f).patterns.patterns:
            for slot in pattern.slots:
                if slot.capacity and slot.capacity.max_chars:
                    by_role.setdefault(slot.role, []).append(slot.capacity.max_chars)
    for role, caps in sorted(by_role.items(), key=lambda kv: -len(kv[1])):
        q1, med, q3 = _quartiles(caps)
        print(f"| {role} | {len(caps)} | {q1} | {med} | {q3} |")

    print("\n### Колода из этого входа по шаблонам\n")
    print("| Шаблон | Слайдов | Не размещено | Заливок | Сверх ёмкости | Медиана знаков в заливке |")
    print("|---|---|---|---|---|---|")
    empty = 0
    total = 0
    for f in samples():
        total += 1
        name = os.path.basename(f)
        try:
            a = analyze_template(f)
            p = plan_deck(doc, a.patterns, a.design_system.source.sha256)
        except Exception as exc:  # noqa: BLE001 — падение обязано быть видно в таблице
            print(f"| `{name}` | ПАДЕНИЕ {type(exc).__name__} | — | — | — | — |")
            empty += 1
            continue
        fills = [fill for s in p.slides for fill in s.fills]
        over = sum(1 for fill in fills if fill.over_capacity)
        lengths = [len(fill.text) for fill in fills if getattr(fill, "text", None)]
        median = _quartiles(lengths)[1] if lengths else "—"
        empty += not p.slides
        print(f"| `{name}` | {len(p.slides)} | {len(p.unplaced)} | {len(fills)} "
              f"| {over} | {median} |")
    print(f"\n**Итог: пустых колод {empty} из {total}.**")


def report_volume(content: str, slides: str | None) -> None:
    """Замеры по `Z-35`: что делает цель по объёму. Сравнивает «без цели» и
    «с целью» на одном и том же входе и шаблоне."""
    from mimeo.cli import parse_slides

    target = parse_slides(slides or "10-15")
    doc_free = load_content(content)
    doc_aim = load_content(content, target=target)
    print(f"### Вход `{os.path.basename(content)}`, цель {target[0]}-{target[1]}")
    print()
    print(f"Разделов: без цели {len(doc_free.sections)}, с целью {len(doc_aim.sections)}")
    print()
    print("| Шаблон | Раскладок | Слайдов без цели | Слайдов с целью | Сверх ёмкости без/с |")
    print("|---|---|---|---|---|")
    for f in samples():
        a = analyze_template(f)
        sha = a.design_system.source.sha256
        free = plan_deck(doc_free, a.patterns, sha)
        aim = plan_deck(doc_aim, a.patterns, sha, target=target)

        def over(p):
            return sum(1 for s in p.slides for fill in s.fills if fill.over_capacity)

        print(f"| `{os.path.basename(f)[:44]}` | {len(a.patterns.patterns)} "
              f"| {len(free.slides)} | {len(aim.slides)} | {over(free)}/{over(aim)} |")


#: Отложенная часть корпуса (`ADR-0019`). Настраиваемся на своих одиннадцати
#: шаблонах, проверяемся на трёх выданных — их выбирали не мы, и это самое
#: близкое к финальной ситуации, что у нас есть. Список зафиксирован здесь, а не
#: живёт в голове: иначе его незаметно подвинут в удобную сторону.
HOLDOUT = (
    "2_Датасет VK Tech шаблон.pptx",
    "2_Датасет VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
    "2_Датасет Шаблон презентации VK Education.pptx",
)


def report_quality() -> None:
    """Весы колоды, дешёвый ярус: базовая линия по корпусу (`PLAN-2.5`).

    Столбцы — ступени сравнения по старшинству: потери, поломки, разнообразие,
    «на донышке». Меньше лучше везде, кроме разнообразия.
    """
    from mimeo.plan.quality import score_deck

    inputs = {"demo": load_content(CONTENT),
              "prose": load_content("examples/content-prose.md")}
    print("| Шаблон | Часть | Вход | Слайдов | Заливок | Потеряно | Сломано | "
          "Тесно | Раскладок | На донышке |")
    print("|---|---|---|---|---|---|---|---|---|---|")
    totals = {}
    for f in samples():
        name = os.path.basename(f)
        part = "отложенная" if name in HOLDOUT else "настроечная"
        a = analyze_template(f)
        sha = a.design_system.source.sha256
        for label, doc in inputs.items():
            plan = plan_deck(doc, a.patterns, sha)
            sc = score_deck(plan, a.patterns)
            acc = totals.setdefault((part, label), [0, 0, 0, 0, 0, 0])
            acc[0] += sc.slides; acc[1] += sc.fills; acc[2] += sc.lost
            acc[3] += sc.broken; acc[4] += sc.tight; acc[5] += sc.thin
            print(f"| `{name[:42]}` | {part} | {label} | {sc.slides} | {sc.fills} "
                  f"| {sc.lost} | {sc.broken} | {sc.tight} | {sc.layouts} | {sc.thin} |")
    print()
    print("| Часть | Вход | Слайдов | Заливок | Потеряно | Сломано | Тесно | На донышке |")
    print("|---|---|---|---|---|---|---|---|")
    for (part, label), a in sorted(totals.items()):
        print(f"| {part} | {label} | {a[0]} | {a[1]} | {a[2]} | {a[3]} | {a[4]} | {a[5]} |")


def main() -> int:
    global TEMPLATES
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "stage",
        choices=("analyze", "plan", "compose", "prose", "volume", "quality", "llm",
                 "slots", "verify"),
    )
    parser.add_argument("--content", default="examples/content-prose.md",
                        help="вход для stage=prose (по умолчанию прозаический пример)")
    parser.add_argument("--slides", help="цель для stage=volume (по умолчанию 10-15)")
    parser.add_argument("--raw", action="store_true",
                        help="stage=prose: без сегментации прозы, состояние до Z-25")
    parser.add_argument("--variants", type=int, default=1,
                        help="stage=verify: сколько вариантов вёрстки на шаблон; "
                             "3 на tz/templates даёт девять сдаточных колод")
    parser.add_argument("--templates", default=TEMPLATES,
                        help="каталог с шаблонами (по умолчанию samples)")
    args = parser.parse_args()
    TEMPLATES = args.templates
    if args.stage == "prose":
        report_prose(args.content, args.raw)
        return 0
    if args.stage == "volume":
        report_volume(args.content, args.slides)
        return 0
    if args.stage == "quality":
        report_quality()
        return 0
    if args.stage == "llm":
        report_llm(args.content)
        return 0
    if args.stage == "slots":
        report_slots(args.content)
        return 0
    if args.stage == "verify":
        report_verify(args.content, args.slides, args.variants)
        return 0
    {"analyze": report_analyze, "plan": report_plan, "compose": report_compose}[args.stage]()
    return 0


if __name__ == "__main__":
    sys.exit(main())
